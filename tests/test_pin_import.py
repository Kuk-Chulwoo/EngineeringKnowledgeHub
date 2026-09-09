import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app.config import Settings
from backend.app.main import create_app
from backend.app.pin_import.parser import MAX_PIN_ROWS
from backend.app.pin_import.service import PinImportService


@pytest.fixture
def client(tmp_path):
    settings = Settings(tmp_path / "hub.sqlite3", tmp_path / "files", 4 * 1024 * 1024)
    with TestClient(create_app(settings)) as test_client:
        yield test_client


def register(client):
    response = client.post("/api/v1/components", json={
        "manufacturer": "Acme", "part_number": "IMPORT", "description": "Import test",
        "category": "Test", "package": "QFN",
    })
    assert response.status_code == 201
    return response.json()["id"]


def xlsx(rows):
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def preview(client, component_id, filename, content):
    return client.post(
        f"/api/v1/components/{component_id}/pin-import/preview",
        files={"file": (filename, content)},
    )


def test_csv_and_xlsx_templates(client):
    csv_response = client.get("/api/v1/pin-import/template?format=csv")
    assert csv_response.status_code == 200
    assert csv_response.content.startswith(b"\xef\xbb\xbfPin Number,Pin Name\r\n")
    assert "pin_template.csv" in csv_response.headers["content-disposition"]

    xlsx_response = client.get("/api/v1/pin-import/template?format=xlsx")
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    assert list(workbook.active.values) == [("Pin Number", "Pin Name")]
    workbook.close()
    assert client.get("/api/v1/pin-import/template?format=xls").status_code == 422


def test_csv_preview_normalizes_bom_and_skips_blank_rows_without_writing(client):
    component_id = register(client)
    content = "\ufeff pin_number , PIN_NAME \r\n1,RESET_N\r\nA1,DIO_0\r\nEP,GND\r\n,\r\n".encode()
    response = preview(client, component_id, "pins.csv", content)
    assert response.status_code == 200
    result = response.json()
    assert result["valid"] is True
    assert result["count"] == 3
    assert result["pins"] == [
        {"pin_number": "1", "pin_name": "RESET_N"},
        {"pin_number": "A1", "pin_name": "DIO_0"},
        {"pin_number": "EP", "pin_name": "GND"},
    ]
    assert client.get(f"/api/v1/components/{component_id}/pins").json()["count"] == 0


def test_csv_validation_errors_are_specific_and_do_not_replace_pins(client):
    component_id = register(client)
    endpoint = f"/api/v1/components/{component_id}/pins"
    client.put(endpoint, json={"source_type": "MANUAL", "pins": [
        {"pin_number": "OLD", "pin_name": "KEEP"}
    ]})
    cases = [
        (b"Other,Pin Name\n1,A", "MISSING_HEADER"),
        (b"Pin Number,pin_number,Pin Name\n1,1,A", "DUPLICATE_HEADER"),
        (b"Pin Number,Pin Name\nEP,A\nep,B", "DUPLICATE_PIN_NUMBER"),
        (b"Pin Number,Pin Name\n,A", "BLANK_PIN_NUMBER"),
        (b"Pin Number,Pin Name\n1,", "BLANK_PIN_NAME"),
        (b"Pin Number,Pin Name\n,\n", "NO_USABLE_ROWS"),
        (b'Pin Number,Pin Name\n"unterminated', "MALFORMED_FILE"),
        (b"\xff\xfe", "MALFORMED_FILE"),
    ]
    for content, code in cases:
        result = preview(client, component_id, "pins.csv", content).json()
        assert result["valid"] is False
        assert result["errors"][0]["code"] == code
    assert client.get(endpoint).json()["pins"][0]["pin_number"] == "OLD"


def test_xlsx_preview_supports_numeric_alphanumeric_ep_and_blank_rows(client):
    component_id = register(client)
    content = xlsx([
        ["Pin Number", "Pin Name"], [1, "RESET_N"], ["B12", "GPIO0"],
        ["EP", "GND"], [None, None], [1.5, "FRACTIONAL_ID"],
    ])
    result = preview(client, component_id, "pins.xlsx", content).json()
    assert result["valid"] is True
    assert [pin["pin_number"] for pin in result["pins"]] == ["1", "B12", "EP", "1.5"]


def test_xlsx_validation_and_malformed_or_unsupported_files(client):
    component_id = register(client)
    for rows, code in [
        ([["Pin Name"], ["A"]], "MISSING_HEADER"),
        ([["Pin Number", "Pin Name"], ["EP", "A"], ["ep", "B"]],
         "DUPLICATE_PIN_NUMBER"),
        ([["Pin Number", "Pin Name"], [1, None]], "BLANK_PIN_NAME"),
    ]:
        result = preview(client, component_id, "pins.xlsx", xlsx(rows)).json()
        assert result["valid"] is False
        assert result["errors"][0]["code"] == code
    malformed = preview(client, component_id, "pins.xlsx", b"not a zip").json()
    assert malformed["errors"][0]["code"] == "MALFORMED_FILE"
    unsupported = preview(client, component_id, "pins.xls", b"legacy").json()
    assert unsupported["errors"][0]["code"] == "UNSUPPORTED_FILE_TYPE"
    assert preview(client, 999, "pins.csv", b"Pin Number,Pin Name\n1,A").status_code == 404


def test_excessive_rows_and_oversized_content_are_rejected():
    service = PinImportService()
    too_many = "Pin Number,Pin Name\n" + "\n".join(
        f"{index},P{index}" for index in range(MAX_PIN_ROWS + 1)
    )
    assert service.preview("pins.csv", too_many.encode())["errors"][0]["code"] == "TOO_MANY_ROWS"
    oversized = service.preview("pins.csv", b"x" * (2 * 1024 * 1024 + 1))
    assert oversized["errors"][0]["code"] == "FILE_TOO_LARGE"


def test_confirm_uses_existing_canonical_put_and_user_import_source(client):
    component_id = register(client)
    endpoint = f"/api/v1/components/{component_id}/pins"
    client.put(endpoint, json={"source_type": "MANUAL", "pins": [
        {"pin_number": "OLD", "pin_name": "OLD_NAME"}
    ]})
    parsed = preview(
        client, component_id, "new.csv", b"Pin Number,Pin Name\n2,GPIO0\n1,RESET_N"
    ).json()
    response = client.put(endpoint, json={"source_type": "USER_IMPORT", "pins": parsed["pins"]})
    assert response.status_code == 200
    assert response.json()["source_type"] == "USER_IMPORT"
    assert [pin["pin_number"] for pin in response.json()["pins"]] == ["1", "2"]
