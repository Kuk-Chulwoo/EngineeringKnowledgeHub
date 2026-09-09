import csv
import io
import math
import zipfile
from collections.abc import Iterable
from itertools import islice
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_PIN_ROWS = 5000
MAX_XLSX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_XLSX_MEMBERS = 100

HEADERS = {"pin number": "pin_number", "pin_number": "pin_number",
           "pin name": "pin_name", "pin_name": "pin_name"}


class ParseFailure(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message


def safe_filename(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).name[:255]


def detect_format(filename: str) -> str:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ParseFailure("UNSUPPORTED_FILE_TYPE", "Only .csv and .xlsx files are supported")
    return suffix[1:]


def parse_csv(content: bytes) -> list[tuple[int, Any, Any]]:
    try:
        text = content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text, newline=""), strict=True))
    except (UnicodeDecodeError, csv.Error) as error:
        raise ParseFailure("MALFORMED_FILE", "CSV file is not valid UTF-8 CSV") from error
    return _select_columns(rows)


def _check_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise ParseFailure("MALFORMED_FILE", "XLSX archive contains too many entries")
            if sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ParseFailure("FILE_TOO_LARGE", "Expanded XLSX exceeds the safety limit")
    except zipfile.BadZipFile as error:
        raise ParseFailure("MALFORMED_FILE", "XLSX file is malformed") from error


def parse_xlsx(content: bytes) -> list[tuple[int, Any, Any]]:
    _check_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            worksheet = workbook.worksheets[0]
            rows = [
                list(row)
                for row in islice(worksheet.iter_rows(values_only=True), MAX_PIN_ROWS + 2)
            ]
        finally:
            workbook.close()
    except Exception as error:
        raise ParseFailure("MALFORMED_FILE", "XLSX file is malformed") from error
    return _select_columns(rows)


def _select_columns(rows: list[list[Any]]) -> list[tuple[int, Any, Any]]:
    if not rows:
        raise ParseFailure("MISSING_HEADER", "Required headers are missing")
    indexes: dict[str, int] = {}
    for index, value in enumerate(rows[0]):
        normalized = str(value).strip().lower() if value is not None else ""
        field = HEADERS.get(normalized)
        if field and field in indexes:
            raise ParseFailure("DUPLICATE_HEADER", f"Duplicate required header: {field}")
        if field:
            indexes[field] = index
    missing = [field for field in ("pin_number", "pin_name") if field not in indexes]
    if missing:
        raise ParseFailure("MISSING_HEADER", "Missing required header: " + ", ".join(missing))
    if len(rows) - 1 > MAX_PIN_ROWS:
        raise ParseFailure("TOO_MANY_ROWS", f"Pin import cannot exceed {MAX_PIN_ROWS} rows")
    selected = []
    for row_number, row in enumerate(rows[1:], start=2):
        selected.append((row_number, _cell(row, indexes["pin_number"]),
                         _cell(row, indexes["pin_name"])))
    return selected


def _cell(row: list[Any], index: int) -> Any:
    return row[index] if index < len(row) else None


def normalized_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return str(value)
    return str(value).strip()


def validate_rows(rows: Iterable[tuple[int, Any, Any]]) -> tuple[list[dict[str, str]], list[dict]]:
    pins: list[dict[str, str]] = []
    errors: list[dict] = []
    seen: dict[str, int] = {}
    for row_number, raw_number, raw_name in rows:
        number, name = normalized_cell(raw_number), normalized_cell(raw_name)
        if not number and not name:
            continue
        if not number:
            errors.append(_issue("BLANK_PIN_NUMBER", row_number, "pin_number", "Pin Number is required"))
            continue
        if not name:
            errors.append(_issue("BLANK_PIN_NAME", row_number, "pin_name", "Pin Name is required"))
            continue
        duplicate_key = number.casefold()
        if duplicate_key in seen:
            errors.append(_issue("DUPLICATE_PIN_NUMBER", row_number, "pin_number",
                                 "Duplicate pin number"))
            continue
        seen[duplicate_key] = row_number
        pins.append({"pin_number": number, "pin_name": name})
    if not pins and not errors:
        errors.append(_issue("NO_USABLE_ROWS", None, None, "No usable pin rows were found"))
    return pins, errors


def _issue(code: str, row: int | None, field: str | None, message: str) -> dict:
    return {"code": code, "row": row, "field": field, "message": message}
